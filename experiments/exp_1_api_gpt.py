"""
This experiment leverages the ChatGPT 4.1 mini and cosine similarity evaluation

"""
import os
import pandas as pd
from openpyxl import load_workbook
from openai import OpenAI
from dotenv import load_dotenv

import sys_config
from tools.similarity import calculate_cosine_similarity
from tools.excel_operator import copy_qa_sheet_as_target_sheet
from tools.data_loader import get_excel_data_path
from tools.summary_recorder import SummaryEntity, append_summary

# init open ai client
load_dotenv()
client = OpenAI(api_key=os.getenv('OPENAI_API_KEY'))


def ask_open_ai(q):
    res = client.responses.create(
        model=config.OPEN_AI_MODEL,
        instructions='',
        input=q,
    )
    return res.output_text


def asking_each_questions(excel_name: str, sheet_name: str):
    excel_path = get_excel_data_path(excel_name)

    # read and deal data
    df = pd.read_excel(excel_path, sheet_name=sheet_name)
    questions_col = df['three_asking_ways']
    answers = []
    for i, q in enumerate(questions_col):
        answer = ask_open_ai(q)
        answers.append(answer)
    df['answers'] = answers

    # similarity
    df['standard_answers'] = df['standard_answers'].ffill()
    similarity_list = calculate_cosine_similarity(
        df['standard_answers'].astype(str).tolist(),
        df['answers'].astype(str).tolist(),
    )
    df['similarity'] = similarity_list

    # read file and column by openpyxl
    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    header = [cell.value for cell in ws[1]]
    answer_column_index = header.index('answers') + 1
    similarity_column_index = header.index('similarity') + 1

    # write answers row by row
    for i, v in enumerate(df['answers'], start=2):
        ws.cell(row=i, column=answer_column_index, value=v)

    # write similarity row by row
    for i, v in enumerate(df['similarity'], start=2):
        ws.cell(row=i, column=similarity_column_index, value=v)

    # use this to save for keeping style of the sheet
    wb.save(excel_path)

    # record the summary
    summary = SummaryEntity(
        experiment_name=excel_name,
        dataset_version='v1',
        book_name='all',
        filter_name='',
        model_name=config.OPEN_AI_MODEL,
        db_collection_name='',
        question_number=len(answers),
    )
    append_summary(summary, similarity_list)


if __name__ == '__main__':
    data_file_name = 'exp_1_api_gpt'
    target_sheet = copy_qa_sheet_as_target_sheet(data_file_name, 'all_book')
    asking_each_questions(data_file_name, target_sheet)



