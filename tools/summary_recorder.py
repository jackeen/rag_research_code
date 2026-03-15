"""
The IO of the result of the experiments.
It appends the results into the Excel file.
"""
from datetime import datetime
import numpy as np
from openpyxl import load_workbook
from dataclasses import dataclass, fields
from .excel_operator import get_excel_data_path


@dataclass
class SummaryEntity:
    experiment_name: str
    dataset_version: str
    ingest_file_name: str

    db_collection_name: str
    question_number: int

    model_name: str
    hybrid_retrieving: bool = False
    dense_embedding_model_name: str = ''
    sparse_embedding_model_name: str = ''

    filter_name: str = ''
    filter_version: str = ''
    filter_params: str = ''

    chunking_size: int = 100
    chunking_overlap: int = 0

    average_similarity: float = 0.0
    median_similarity: float = 0.0
    average_bert_score: float = 0.0
    median_bert_score: float = 0.0
    nli_cls_contradiction_percentage: float = 0.0
    nli_cls_entailment_percentage: float = 0.0
    nli_cls_neutral_percentage: float = 0.0

    experiment_time: str | None = None


def append_summary(record: SummaryEntity, similarity_list: list[float]):
    """
    Append a summary entity to the given sheet with experiment name.
    :param record:
    :param similarity_list:
    :return:
    """
    record.average_similarity = round(float(np.mean(similarity_list)), 4)
    record.median_similarity = round(float(np.median(similarity_list)), 4)
    record.experiment_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    summary_path = str(get_excel_data_path('summary'))
    wb = load_workbook(summary_path)
    ws = wb[record.experiment_name]

    # align the sheet columns with fields of dataclass
    sheet_columns = [cell.value for cell in ws[1]]
    record_dict = {f.name: getattr(record, f.name) for f in fields(record)}
    record_row = [record_dict.get(col) for col in sheet_columns]

    # append row and save
    ws.append(record_row)
    wb.save(summary_path)


def test_append_summary():
    summary = SummaryEntity(
        experiment_name='exp_1_api_gpt',
        dataset_version='v1',
        ingest_file_name='all',
        filter_name='None',
        model_name='None',
        db_collection_name='None',
        question_number=2,
    )
    append_summary(summary, [0.3, 0.6])


if __name__ == '__main__':
    # test_append_summary()
    pass
