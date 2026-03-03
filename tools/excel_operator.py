"""
These tools are working for Excel file as the experiment data.
"""
import pandas as pd
from openpyxl import load_workbook
from .data_loader import get_excel_data_path


def copy_qa_sheet_as_target_sheet(file_name: str, sheet_name: str) -> str:
    """

    :param file_name:
    :param sheet_name:
    :return: the copied sheet name
    """
    excel_path = get_excel_data_path(file_name)
    wb = load_workbook(str(excel_path))

    # avoid name collision
    i = 1
    target_name = sheet_name
    while target_name in wb.sheetnames:
        target_name = f'{sheet_name}_{i}'
        i += 1

    # save sheet and return actual sheet name
    source_sheet = wb['QAs']
    new_sheet = wb.copy_worksheet(source_sheet)
    new_sheet.title = target_name
    wb.save(excel_path)
    return target_name


def write_target_sheet_column_cells(
        file_name: str,
        sheet_name: str,
        column_name: str,
        data: list[any],
        offset: int,
):
    """
    Write the target sheet column to the Excel file.
    :param file_name: the target file name
    :param sheet_name: the target sheet name
    :param column_name: the target column name
    :param data: the data list for target cells
    :param offset: the start of rows' offset
    :return:
    """
    excel_path = str(get_excel_data_path(file_name))
    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    header = [cell.value for cell in ws[1]]
    target_column_index = header.index(column_name) + 1
    for i, v in enumerate(data, start=2+offset):
        ws.cell(row=i, column=target_column_index, value=v)

    wb.save(excel_path)
